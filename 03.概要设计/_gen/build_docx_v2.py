# -*- coding: utf-8 -*-
"""把《概要设计说明书 v2.0》Markdown 转成 Word 交付件。
复用 build_docx.py 的排版逻辑（A4 + 页码 + 表格样式 + 图片限宽），只替换源文件与资源路径。
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(r"D:\Desktop\康养系统实训项目\03.概要设计")
sys.path.insert(0, str(BASE))
import build_docx as bd  # noqa: E402

SOURCES = ["概要设计说明书_康养系统_v2.0.md"]
TITLE = "基于 AI 智能体的智能康养系统"
TOC_HEADINGS = ("Table of Contents", "目录", "Contents")


def page_break_paragraph():
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    p.append(r)
    return p


def fix_cover_order(path: Path) -> None:
    """pandoc 把目录放在最前面；按模板顺序改为 封面 → 文件修改控制 → 目录 → 正文"""
    doc = Document(str(path))
    body = doc.element.body
    sdt = body.find(qn("w:sdt"))
    if sdt is None:
        print("  [warn] 未找到目录域，跳过封面顺序调整")
        return
    body.remove(sdt)
    tables = doc.tables
    anchor = tables[1]._tbl if len(tables) > 1 else tables[0]._tbl
    br1 = page_break_paragraph()
    anchor.addnext(br1)
    br1.addnext(sdt)

    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "1 文档概述":
            paragraph._p.addprevious(page_break_paragraph())
            break
    doc.save(str(path))


def main() -> None:
    pandoc = shutil.which("pandoc") or str(bd.PANDOC_FALLBACK)
    for name in SOURCES:
        source = BASE / name
        output = source.with_suffix(".docx")
        with tempfile.TemporaryDirectory(prefix="ky2_docx_") as td:
            raw = Path(td) / "raw.docx"
            subprocess.run(
                [pandoc, str(source), "-o", str(raw),
                 "--from=markdown-smart", "--dpi=300", "--toc", "--toc-depth=3",
                 "--metadata=toc-title=目录",
                 "--resource-path=%s;%s" % (BASE, BASE / "图_v2")],
                cwd=str(BASE), check=True)
            bd.style_document(raw, output)
        fix_cover_order(output)
        print("generated:", output)


if __name__ == "__main__":
    main()


