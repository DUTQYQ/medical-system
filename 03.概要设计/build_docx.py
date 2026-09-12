"""从 Markdown 生成适合评审和打印的 Word 交付件。"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path(__file__).resolve().parent
SOURCES = (
    "概要设计说明书_康养系统_v1.0.md",
    "交付说明_2026-09-12.md",
)
PANDOC_FALLBACK = Path.home() / "AppData/Local/Pandoc/pandoc.exe"


def set_font(run, name: str, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), "D9D9D9")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_fixed_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def column_weights(headers: list[str]) -> list[float]:
    n = len(headers)
    joined = "|".join(headers)
    if n == 2:
        return [0.34, 0.66]
    if n == 3:
        return [0.22, 0.34, 0.44]
    if n == 4 and "字段" in headers and "类型" in headers:
        return [0.22, 0.20, 0.23, 0.35]
    if n == 4 and "版本" in headers:
        return [0.13, 0.17, 0.14, 0.56]
    if n == 4 and "路径" in headers:
        return [0.20, 0.36, 0.24, 0.20]
    if n == 4:
        return [0.15, 0.27, 0.27, 0.31]
    if n == 5 and "数据表" in joined:
        return [0.13, 0.18, 0.18, 0.27, 0.24]
    if n == 5:
        return [0.09, 0.25, 0.30, 0.17, 0.19]
    return [1.0 / n] * n


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(Cm(width_cm).twips)))
    tc_w.set(qn("w:type"), "dxa")


def style_table(table, usable_width_cm: float) -> None:
    if not table.rows:
        return
    headers = [cell.text.strip() for cell in table.rows[0].cells]
    weights = column_weights(headers)
    font_size = 8.0 if len(headers) >= 5 else 8.5 if len(headers) == 4 else 9.0
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_fixed_layout(table)
    set_table_borders(table)
    set_repeat_header(table.rows[0])

    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_width(cell, usable_width_cm * weights[min(column_index, len(weights) - 1)])
            if row_index == 0:
                shade_cell(cell, "D9E2F3")
            elif row_index % 2 == 0:
                shade_cell(cell, "F7F9FC")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    set_font(run, "宋体", font_size, bold=True if row_index == 0 else None)
                    run.font.color.rgb = RGBColor(0, 0, 0)


def remove_rule_paragraphs(document: Document) -> None:
    for paragraph in list(document.paragraphs):
        p_pr = paragraph._p.pPr
        if not paragraph.text.strip() and p_pr is not None and p_pr.find(qn("w:pBdr")) is not None:
            paragraph._element.getparent().remove(paragraph._element)


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.clear()
    run = paragraph.add_run("第 ")
    set_font(run, "宋体", 9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_font(run, "宋体", 9)


def style_document(raw_path: Path, output_path: Path) -> None:
    document = Document(raw_path)
    usable_width_cm = 21.0 - 1.8 - 1.8
    for section in document.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
        add_page_number(section)

    styles = document.styles
    style_specs = {
        "Normal": ("宋体", 10.5, False),
        "Title": ("黑体", 20, True),
        "Heading 1": ("黑体", 16, True),
        "Heading 2": ("黑体", 13.5, True),
        "Heading 3": ("黑体", 11.5, True),
        "Heading 4": ("黑体", 10.5, True),
    }
    for style_name, (font_name, size, bold) in style_specs.items():
        if style_name not in styles:
            continue
        style = styles[style_name]
        style.font.name = font_name
        r_pr = style._element.get_or_add_rPr()
        r_fonts = r_pr.find(qn("w:rFonts"))
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, r_fonts)
        r_fonts.set(qn("w:eastAsia"), font_name)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)

    first_text = next((p for p in document.paragraphs if p.text.strip()), None)
    if first_text is not None:
        first_text.style = styles["Title"]
        first_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
        first_text.paragraph_format.space_after = Pt(12)

    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        paragraph.paragraph_format.widow_control = True
        if style_name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(5)
        elif style_name != "Title":
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.25
        for run in paragraph.runs:
            if style_name.startswith("Heading") or style_name == "Title":
                run.font.color.rgb = RGBColor(0, 0, 0)
                set_font(run, "黑体")

    remove_rule_paragraphs(document)
    for table in document.tables:
        style_table(table, usable_width_cm)

    for shape in document.inline_shapes:
        if shape.width > Cm(usable_width_cm):
            ratio = Cm(usable_width_cm) / shape.width
            shape.width = Cm(usable_width_cm)
            shape.height = int(shape.height * ratio)

    document.core_properties.author = "丘宇乾"
    document.core_properties.title = first_text.text if first_text is not None else output_path.stem
    document.save(output_path)


def main() -> None:
    pandoc = shutil.which("pandoc")
    if pandoc is None and PANDOC_FALLBACK.exists():
        pandoc = str(PANDOC_FALLBACK)
    if pandoc is None:
        raise SystemExit("未找到 pandoc，请先安装后重试")

    for source_name in SOURCES:
        source = BASE / source_name
        output = source.with_suffix(".docx")
        with tempfile.TemporaryDirectory(prefix="kangyang_docx_") as temp_dir:
            raw = Path(temp_dir) / "raw.docx"
            subprocess.run(
                [
                    pandoc,
                    str(source),
                    "-o",
                    str(raw),
                    "--dpi=300",
                    f"--resource-path={BASE};{BASE / '图'}",
                ],
                cwd=BASE,
                check=True,
            )
            style_document(raw, output)
        print(f"generated: {output}")


if __name__ == "__main__":
    main()
