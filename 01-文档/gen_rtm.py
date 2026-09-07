# -*- coding: utf-8 -*-
"""生成康养系统【老师模板版】需求跟踪矩阵 RTM 的 Excel。
三张表：需求跟踪矩阵封面 / 变更履历 / 设计用 RTM（+ 图1~图7 空页供截图）。
42 条需求，按老师模板 大分類(模块)/中分類(子模块)/小分類(功能点)/详细说明 格式。
详细说明含「参照图X-图标Y」+ UI 规格（字号/颜色/动作）。
占位符：担当者=[组员A/B/C]，责任者=组长。
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

OUT = r"D:\Desktop\康养系统实训项目\01-文档\RTM_康养系统_老师模板版.xlsx"

# 42 条需求: (大分類, 中分類, 小分類, 详细说明)
RTM = [
    # ============ 图1 · 认证（登录/注册） ============
    ("认证（登录/注册）", "登录界面", "header",
     "参照图1，图标1，静态显示\"康养健康助手\"，字号22px，加粗，色号#2d6a4f，居中"),
    ("认证（登录/注册）", "登录界面", "form",
     "参照图1，图标2，占位提示\"手机号\"，字号16px，动作：用户输入已注册手机号，仅允许输入数字"),
    ("认证（登录/注册）", "登录界面", "form",
     "参照图1，图标3，占位提示\"密码\"，字号16px，动作：输入登录密码，内容用圆点遮蔽"),
    ("认证（登录/注册）", "登录界面", "actions",
     "参照图1，图标4，文字\"登录\"，字号18px，白字#FFFFFF，背景#2d6a4f，动作：校验手机号与密码非空，成功后按角色（老人/家属/护工/管理员）签发JWT并跳转对应首页，失败提示\"手机号或密码错误！\""),
    ("认证（登录/注册）", "登录界面", "actions",
     "参照图1，图标5，扁平按钮\"没有账号？点击注册\"，无边框，色号#52b788，字号16px，动作：点击切换至注册页"),
    ("认证（登录/注册）", "注册界面", "header+actions",
     "参照图1，图标6，标题\"用户注册\"，字号20px，加粗，色号#2d6a4f；图标7，注册按钮\"注册\"字号18px白字背景#52b788，动作：校验手机号/密码强度，勾选健康数据隐私授权，成功后创建账号并跳登录"),
    # ============ 图2 · 老人首页（适老化） ============
    ("老人首页（适老化）", "header", "页面标题+健康卡片",
     "参照图2，图标1，标题\"健康首页\"，黑体，字号22px，加粗，颜色#2d6a4f，左对齐；图标2，今日健康卡片动态显示血压/心率/血糖，字号20px，正常值色#2d6a4f，超标值红#e74c3c"),
    ("老人首页（适老化）", "nav", "档案入口",
     "参照图2，图标3，大按钮\"我的档案\"，字号20px，加粗，白字，背景#52b788，高度≥48px，动作：点击直达健康档案查看"),
    ("老人首页（适老化）", "nav", "AI咨询入口",
     "参照图2，图标4，大按钮\"问一问 AI\"，字号20px，加粗，白字，背景#2d6a4f，高度≥48px，动作：点击进入AI咨询页"),
    ("老人首页（适老化）", "nav", "预警入口",
     "参照图2，图标5，大按钮\"我的预警\"，字号20px，加粗，白字，背景#e67e22，存在未处理预警时右上角红色角标#e74c3c，动作：点击直达预警中心"),
    ("老人首页（适老化）", "toolbar", "一键朗读",
     "参照图2，图标6，\"🔊 朗读\"按钮，字号16px，边框#2d6a4f，动作：点击调用网页朗读（Web Speech API）播报页面关键文字，再点停止"),
    ("老人首页（适老化）", "toolbar", "关怀提醒",
     "参照图2，图标15，登录后弹出用药/健康打卡提醒卡片，字号18px，背景#f0faf3浅绿，动作：点击\"知道了\"关闭，每日仅提醒一次"),
    # ============ 图2 · 健康档案 ============
    ("健康档案", "档案列表", "档案切换与查看",
     "参照图2，图标7，以卡片列表展示本账号多份档案（姓名/年龄/照片/慢病标签），字号18px，加粗，动作：点击卡片即切换并查看对应档案详情"),
    ("健康档案", "档案表单", "创建档案",
     "参照图2，图标8，按钮\"新建档案\"，白字，背景#52b788，字号18px，动作：弹出表单填写姓名/年龄/病史/慢病标签，必填校验通过后保存并生成新档案"),
    ("健康档案", "档案表单", "编辑与删除档案",
     "参照图2，图标9，编辑按钮\"保存\"白字背景#2d6a4f；图标10，删除按钮文字红色#e74c3c，动作：编辑后保存即时生效；删除前弹出二次确认，确认后删除并解除绑定"),
    # ============ 图2 · 指标录入 ============
    ("指标录入", "录入表单", "血压录入",
     "参照图2，图标11，标签\"血压\"字号18px，输入收缩/舒张压数值，动作：范围校验（收缩90~200 / 舒张60~130），越界红色#e74c3c提示"),
    ("指标录入", "录入表单", "心率与血糖录入",
     "参照图2，图标12，标签\"心率 / 血糖\"字号18px，动作：录入数值，心率40~200、血糖2.0~20.0 范围校验，越界提示"),
    ("指标录入", "录入表单", "保存录入+趋势图",
     "参照图2，图标13，按钮\"保存指标\"白字背景#2d6a4f字号18px，动作：写入统一健康记录表并触发预警检测；图14，ECharts近30天趋势折线，正常点色#2d6a4f，超阈值点红#e74c3c高亮，支持切换指标类型"),
    # ============ 图3 · AI咨询 ============
    ("AI咨询", "header", "页面标题+档案选择",
     "参照图3，图标1，标题\"AI健康咨询\"，黑体，字号20px，加粗，颜色#2d6a4f，左对齐；图标2，下拉框切换咨询档案，字号16px，动作：选中档案后咨询基于该档案数据"),
    ("AI咨询", "chat", "问答消息列表",
     "参照图3，图标3，问答气泡列表：用户气泡右对齐背景#52b788白字，AI气泡左对齐浅绿背景#f0faf3深字，字号18px，动作：新消息自动滚动到底部，AI长文分段分点展示"),
    ("AI咨询", "chat", "预设问题按钮",
     "参照图3，图标4，\"血压高怎么办\"\"用药时间\"等预设按钮，字号16px，边框#2d6a4f，动作：免打字即点即问"),
    ("AI咨询", "actions", "提问输入框",
     "参照图3，图标5，占位提示\"请输入您的健康问题…\"，字号18px，动作：输入健康问题，内容为空时拦截"),
    ("AI咨询", "actions", "发送与智能路由",
     "参照图3，图标6，按钮\"发送\"，字号18px，白字背景#2d6a4f，动作：后端组装（档案摘要+问题）调LLM API，按意图路由（普通/数据/异常/紧急）返回建议，支持流式输出"),
    ("AI咨询", "history", "历史记录与免责声明",
     "参照图3，图标7，\"历史问答\"标题字号16px加粗，动作：按档案查看历史咨询并可回看；所有页面底部固定声明\"内容仅供参考，不构成医疗诊断\"，字号14px，灰色#999"),
    # ============ 图4 · 预警中心 ============
    ("预警中心", "预警规则", "主动预警与阈值引擎",
     "参照图4，图标1，标题\"预警中心\"，黑体，字号20px，加粗，颜色#e74c3c，左对齐；动作：新健康数据写入后自动做阈值检测，命中即触发预警（阈值入sys_config，不写死）"),
    ("预警中心", "预警规则", "风险分级与AI摘要",
     "参照图4，图标2，风险等级标识 正常/轻度/较高/高风险（0~3级），正常绿#2d6a4f、轻度橙#e67e22、较高橙红#e74c3c、高风险红#c0392b；AI结合历史趋势生成个性化预警摘要文字"),
    ("预警中心", "预警列表", "预警列表与详情",
     "参照图4，图标3，预警卡片列表展示指标值/时间/风险等级，字号16px；图标4，点击进入详情展示指标值、阈值、近30天趋势，处理状态标签：待处理(灰#999)/已查看(蓝)/已处理(绿)"),
    ("预警中心", "预警通知", "分级推送",
     "参照图4，图标5，预警自动推送：老人首页提醒+角标，已绑定家属端同步通知，高风险升级通知护工后台；图标6，通知中心汇总所有预警，已读/未读状态，未读红点#e74c3c"),
    ("预警中心", "预警处理", "处理闭环",
     "参照图4，图标7，状态流转 待处理→已查看→已处理；家属/护工可登记处理结论与操作日志留痕，动作：点击\"标记已处理\"弹出表单填写处理结果"),
    ("预警中心", "预警通知", "频率控制",
     "参照图4，图标8，同一异常预警冷静期/频控，避免重复轰炸式通知，动作：短时间内同源异常合并去重"),
    # ============ 图5 · 家属端 ============
    ("家属端", "代管", "绑定与代管老人",
     "参照图5，图标1，标题\"家属端\"，黑体，字号20px，加粗，颜色#2d6a4f，左对齐；图标2，\"添加老人\"按钮白字背景#52b788字号18px，动作：通过手机号/邀请码绑定并代管多位老人，可解绑，解绑后权限立即失效"),
    ("家属端", "健康看板", "代管老人概况",
     "参照图5，图标3，以卡片列表展示所有代管老人，含头衔/今日关键指标/最近预警数量，字号16px，正常绿、有预警红角标，动作：点击切换查看该老人健康概况与档案"),
    ("家属端", "健康看板", "看趋势与代录数据",
     "参照图5，图标4，\"健康趋势\"图表入口，ECharts展示代管老人近30天趋势；图标5，\"代录数据\"按钮白字背景#2d6a4f，动作：在家属端为老人录入血压/心率等健康指标"),
    ("家属端", "预警", "接收预警与代管咨询",
     "参照图5，图标6，接收并查看代管老人全部预警，高优先级置顶红#e74c3c；图标7，\"代问 AI\"按钮白字背景#2d6a4f，动作：以所选老人档案发起AI咨询"),
    ("家属端", "预警", "确认处理与通知中心",
     "参照图5，图标8，按钮\"确认已处理/已联系老人\"，字号16px，白字背景#52b788，动作：登记处理状态供追溯；图标9，家属端通知中心汇总所有老人预警，已读/未读状态"),
    # ============ 图6 · 护工端 ============
    ("护工端", "老人工作台", "负责老人列表与状态标识",
     "参照图6，图标1，标题\"护工端\"，黑体，字号20px，加粗，颜色#2d6a4f，左对齐；图标2，负责老人列表按状态分色标识：正常绿#2d6a4f/关注橙#e67e22/异常红#e74c3c，高风险老人置顶，动作：点击进入详情"),
    ("护工端", "老人工作台", "今日待办与筛选",
     "参照图6，图标3，工作台汇总负责老人今日待办（未处理预警/健康打卡），字号16px，异常优先排序；图标4，支持按状态/时段筛选老人"),
    ("护工端", "预警处理", "查看与处理预警",
     "参照图6，图标5，查看负责老人预警列表与详情（指标值/阈值/趋势），字号16px，处理状态标签；图标6，\"登记处理结果\"按钮白字背景#2d6a4f，动作：填写处理结论并留痕，供追溯"),
    # ============ 图7 · 管理员后台 ============
    ("管理员后台", "用户管理", "用户与角色管理",
     "参照图7，图标1，标题\"系统后台\"，黑体，字号20px，加粗，颜色#2d6a4f，左对齐；图标2，用户列表含角色/状态，支持启用/禁用，动作：禁用后该用户不可登录；图标3，角色分配与权限维护"),
    ("管理员后台", "规则配置", "预警阈值配置",
     "参照图7，图标4，\"预警规则\"配置表单，各指标阈值与风险等级（入sys_config，不写死），字号16px，动作：修改后预警逻辑即时生效"),
    ("管理员后台", "知识库", "知识库维护",
     "参照图7，图标5，\"知识库管理\"，支持常见问答/健康文章增删改（仅管理员可改他人只读），字号16px，动作：条目更新后重建向量索引供AI咨询RAG检索"),
    ("管理员后台", "统计", "数据统计与导出",
     "参照图7，图标6，统计看板展示用户数/咨询次数/预警发生与处理数，ECharts图表，字号16px；动作：支持按日/周/月筛选并导出报表"),
]

def style_header(cell, bold=True, size=11, color="FFFFFF", fill="2D6A4F"):
    cell.font = Font(bold=bold, size=size, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def style_cell(cell, size=11, bold=False, halign="left", wrap=True):
    cell.font = Font(bold=bold, size=size)
    cell.alignment = Alignment(horizontal=halign, vertical="center", wrap_text=wrap)

thin = Side(style="thin", color="999999")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ---------- Sheet1: 需求跟踪矩阵封面 ----------
ws = wb.active
ws.title = "需求跟踪矩阵封面"
ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 14
ws.column_dimensions["E"].width = 14
ws.column_dimensions["F"].width = 14
ws.column_dimensions["G"].width = 14
ws.column_dimensions["H"].width = 14

cover_rows = [
    ["密级：机密", "", "", "", "", "第1版", "", ""],
    ["分册名称：无", "", "", "", "", "第 1 册/共 1 册", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "基于 AI 智能体的康养系统", "", "", "", "", ""],
    ["", "", "需求跟踪矩阵", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["实习实训中心", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
    ["总页数", "正文", "", "附录", "", "生效日期", "", ""],
    ["", "编制：", "", "", "", "审批：", "", ""],
    ["", "", "", "", "", "", "", ""],
]
for r, row in enumerate(cover_rows, start=1):
    for c, val in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=val if val != "" else None)
        cell.border = border
        if r == 11:  # 项目名
            style_cell(cell, 20, bold=True, halign="center")
        elif r == 12:  # 需求跟踪矩阵
            style_cell(cell, 24, bold=True, halign="center")
        elif r == 1:
            style_cell(cell, 11)
        elif r == 19:
            style_cell(cell, 14, bold=True)
        else:
            style_cell(cell, 11)

# ---------- Sheet2: 变更履历 ----------
ws2 = wb.create_sheet("变更履历")
ws2.column_dimensions["A"].width = 6
ws2.column_dimensions["B"].width = 12
ws2.column_dimensions["C"].width = 40
ws2.column_dimensions["D"].width = 10
ws2.column_dimensions["E"].width = 10
ws2.column_dimensions["F"].width = 10
ws2.column_dimensions["G"].width = 14
for L in ["H", "I", "J", "K", "L", "M", "N", "O", "P"]:
    ws2.column_dimensions[L].width = 14

ws2.cell(row=1, column=1, value="变更履历")
ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=16)
style_header(ws2.cell(row=1, column=1), size=14, fill="2D6A4F")

headers2 = ["序号", "日期", "变更内容", "RTM No.", "需求版本", "变更人",
            "初始需求个数", "新增需求个数", "删除需求个数", "修改需求个数",
            "总需求个数", "需求变更比", "对规模的影响", "对工作量的影响", "对进度的影响"]
for c, h in enumerate(headers2, start=1):
    cell = ws2.cell(row=2, column=c, value=h)
    style_header(cell, fill="52B788", size=10)
    cell.border = border

# 17 行记录（row3~row19）
hist_1 = ["1", "2026-09-07", "首次编制：按老师模板（旧衣回收碳减排追踪与积分商城系统·需求跟踪矩阵）格式，将需求分析说明书功能需求整理为 42 条 RTM，按 大分類/中分類/小分類 归类，图号=原型评审图号", "设计用RTM", "v1.0", "组长", "42", "0", "0", "0", "42", "0%", "无", "无", "无"]
for r in range(3, 20):  # rows 3..19 => 17 行
    n = r - 2
    if n == 1:
        vals = hist_1
    else:
        vals = [str(n), "", "", "", "", "", "", "", "", "", "", "", "", "", ""]
    for c, v in enumerate(vals, start=1):
        cell = ws2.cell(row=r, column=c, value=v if v else None)
        style_cell(cell, halign="center")
        cell.border = border

ws2.cell(row=20, column=1, value="SUM")
for c in range(1, 16):
    cell = ws2.cell(row=20, column=c, value="")
    if c == 1:
        cell.value = "SUM"
    style_header(cell, fill="7F8C8D")
    cell.border = border

# ---------- Sheet3: 设计用RTM ----------
ws3 = wb.create_sheet("设计用RTM")
ws3.column_dimensions["A"].width = 6
ws3.column_dimensions["B"].width = 16
ws3.column_dimensions["C"].width = 12
ws3.column_dimensions["D"].width = 14
ws3.column_dimensions["E"].width = 70
ws3.column_dimensions["F"].width = 7
ws3.column_dimensions["G"].width = 7
ws3.column_dimensions["H"].width = 7
ws3.column_dimensions["I"].width = 12
ws3.column_dimensions["J"].width = 10

ws3.cell(row=1, column=1, value="需求跟踪矩阵（RTM）")
ws3.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
style_header(ws3.cell(row=1, column=1), size=14, fill="2D6A4F")

ws3.cell(row=2, column=1, value="项目名称：")
ws3.cell(row=2, column=2, value="基于 AI 智能体的康养系统")
ws3.merge_cells(start_row=2, start_column=2, end_row=2, end_column=5)
ws3.cell(row=2, column=6, value="PM:吴剑光")
ws3.merge_cells(start_row=2, start_column=6, end_row=2, end_column=8)
ws3.cell(row=2, column=9, value="担当:具体参与设计的成员名单")
ws3.merge_cells(start_row=2, start_column=9, end_row=2, end_column=10)
for c in range(1, 11):
    style_cell(ws3.cell(row=2, column=c), 11)
ws3.merge_cells(start_row=2, start_column=1, end_row=2, end_column=1)

# 表头
hdrs = ["Ｎｏ.", "大分類（模块）", "中分類（子模块）", "小分類（功能点）", "详细说明", "PD", "COD", "UT", "担当者", "责任者"]
for c, h in enumerate(hdrs, start=1):
    cell = ws3.cell(row=3, column=c, value=h)
    style_header(cell, fill="52B788")
    cell.border = border

note = "○：完成（通过评审或测试）   △：进行中   ×：未着手   N/A：不适用（没有此项活动）"
ws3.cell(row=4, column=1, value=note)
ws3.merge_cells(start_row=4, start_column=1, end_row=4, end_column=10)
style_cell(ws3.cell(row=4, column=1), 10, halign="left")

# 数据：每一条需求占 1 行，但同名相邻的大分類/中分類留空（同老师模板的合并视觉）
r = 5
prev_big = None
prev_mid = None
for i, (big, mid, small, detail) in enumerate(RTM, start=1):
    b_v = big if big != prev_big else ""
    m_v = mid if (big != prev_big or mid != prev_mid) else ""
    cells = [
        (1, str(i)),
        (2, b_v),
        (3, m_v),
        (4, small),
        (5, detail),
        (6, ""), (7, ""), (8, ""),
        (9, "[组员A/B/C]"),
        (10, "组长"),
    ]
    for c, v in cells:
        cell = ws3.cell(row=r, column=c, value=v if v else None)
        halign = "left" if c in (4, 5) else "center"
        style_cell(cell, 10.5, halign=halign)
        cell.border = border
        if c in (3,):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        if c == 1:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        if c == 10:
            cell.alignment = Alignment(horizontal="center", vertical="center")
    prev_big = big
    prev_mid = mid
    r += 1

# ---------- 图1~图7 空页（供截图） ----------
for i in range(1, 8):
    wb.create_sheet(f"图{i}")
    ws_img = wb[f"图{i}"]
    ws_img.cell(row=1, column=1, value=f"图{i}（此处放原型截图：")
    fig_name = {1: "登录/注册", 2: "老人首页", 3: "AI咨询", 4: "预警中心",
                5: "家属端", 6: "护工端", 7: "管理员后台"}[i]
    ws_img.cell(row=2, column=1, value=f"图{i} = {fig_name}")
    ws_img.column_dimensions["A"].width = 40
    style_cell(ws_img.cell(row=1, column=1), 12, bold=True)
    style_cell(ws_img.cell(row=2, column=1), 11)

wb.save(OUT)
print("SAVED:", OUT)
print("RTM 需求条数:", len(RTM))
